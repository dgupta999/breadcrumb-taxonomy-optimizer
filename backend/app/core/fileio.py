"""CSV/XLSX reading, column-mapping suggestions and row import."""
from __future__ import annotations

import csv
import io
import re
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import delete
from sqlalchemy.orm import Session

from ..models import SRP, Category, Synonym
from .normalize import norm_key

# target field -> (label, required, header aliases)
TAXONOMY_FIELDS: dict[str, tuple[str, bool, list[str]]] = {
    "external_id": ("Category ID", False, ["category id", "cat id", "id", "category_id", "分類id", "類別id"]),
    "url": ("Category URL", True, ["category url", "url", "link", "page url", "網址", "分類網址"]),
    "name_cn": ("Category Name CN", False, ["category name cn", "category name", "name cn", "name", "分類名稱", "類別名稱", "名稱", "category name (cn)"]),
    "name_en": ("Category Name EN", False, ["category name en", "name en", "english name", "category name (en)", "英文名稱"]),
    "h1_cn": ("H1 CN", False, ["h1 cn", "h1", "h1 (cn)"]),
    "h1_en": ("H1 EN", False, ["h1 en", "h1 (en)"]),
    "breadcrumb_cn": ("Breadcrumb CN", True, ["breadcrumb cn", "breadcrumb", "breadcrumbs", "麵包屑", "breadcrumb (cn)", "path"]),
    "breadcrumb_en": ("Breadcrumb EN", False, ["breadcrumb en", "breadcrumb (en)", "english breadcrumb"]),
    "parent_ref": ("Parent Category", False, ["parent category", "parent", "parent name", "上層分類"]),
    "parent_external_id": ("Parent ID", False, ["parent id", "parent_id", "parent category id"]),
    "product_count": ("Product Count", False, ["product count", "products", "product_count", "商品數", "# products"]),
    "status": ("Status", False, ["status", "狀態"]),
}

SRP_FIELDS: dict[str, tuple[str, bool, list[str]]] = {
    "keyword": ("SRP Keyword", True, ["srp keyword", "keyword", "query", "search term", "關鍵字", "kw"]),
    "url": ("SRP URL", True, ["srp url", "url", "page url", "landing page", "網址"]),
    "h1": ("H1", False, ["h1", "title", "h1 cn"]),
    "keyword_en": ("English Keyword", False, ["english keyword", "keyword en", "en keyword", "english", "translation"]),
    "search_volume": ("Search Volume", False, ["search volume", "volume", "sv", "avg monthly searches", "搜尋量"]),
    "product_count": ("Product Count", False, ["product count", "products", "商品數", "results"]),
    "status": ("Status", False, ["status", "狀態"]),
}

SYNONYM_FIELDS: dict[str, tuple[str, bool, list[str]]] = {
    "alias": ("Alias", True, ["alias", "synonym", "term"]),
    "canonical": ("Canonical Concept", True, ["canonical concept", "canonical", "concept"]),
    "preferred_category_name": ("Preferred Category Name", False, ["preferred category name", "category name", "preferred category"]),
    "preferred_node_id": ("Preferred Category ID", False, ["preferred category id", "category id"]),
    "notes": ("Notes", False, ["notes", "note"]),
    "approved_by": ("Approved By", False, ["approved by", "approver"]),
}

COMPETITOR_CRAWL_FIELDS: dict[str, tuple[str, bool, list[str]]] = {
    "url": ("URL", True, ["address", "url", "page url", "網址", "link"]),
    "breadcrumb": ("Breadcrumb", False, ["breadcrumb", "breadcrumbs", "breadcrumb 1", "麵包屑", "breadcrumb text", "breadcrumb path"]),
    "title": ("Title", False, ["title 1", "title", "page title", "meta title"]),
    "h1": ("H1", False, ["h1-1", "h1 1", "h1", "heading"]),
    "status_code": ("Status Code", False, ["status code", "status", "http status"]),
    "canonical": ("Canonical", False, ["canonical link element 1", "canonical", "canonical url"]),
    "indexability": ("Indexability", False, ["indexability", "indexable"]),
    "page_type": ("Page Type", False, ["page type", "type", "template"]),
    "parent_url": ("Parent URL", False, ["parent url", "source", "referrer", "from"]),
    "anchor_text": ("Anchor Text", False, ["anchor text", "anchor", "link text"]),
    "depth": ("Crawl Depth", False, ["crawl depth", "depth", "level"]),
    "product_count": ("Product Count", False, ["product count", "products", "商品數"]),
}

FIELDS = {"TAXONOMY": TAXONOMY_FIELDS, "SRP": SRP_FIELDS, "SYNONYMS": SYNONYM_FIELDS, "COMPETITOR_CRAWL": COMPETITOR_CRAWL_FIELDS}

_BC_COL = re.compile(r"^\s*(breadcrumb|bread crumb|麵包屑)[\s_\-]*(\d+)\s*$", re.I)


def breadcrumb_columns(columns: list[str]) -> list[str]:
    """Numbered breadcrumb columns from crawler exports (Breadcrumb 1, Breadcrumb 2, …), in order."""
    hits = [(int(m.group(2)), c) for c in columns if (m := _BC_COL.match(str(c)))]
    return [c for _, c in sorted(hits)]


def read_table(path: str | Path, max_rows: int | None = None) -> tuple[list[str], list[dict]]:
    p = Path(path)
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        wb = load_workbook(p, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = None
        for row in it:
            if row and any(c not in (None, "") for c in row):
                header = [str(c).strip() if c is not None else f"Column {i+1}" for i, c in enumerate(row)]
                break
        if header is None:
            return [], []
        rows = []
        for row in it:
            if not row or all(c in (None, "") for c in row):
                continue
            rows.append({header[i]: row[i] for i in range(min(len(header), len(row)))})
            if max_rows and len(rows) >= max_rows:
                break
        wb.close()
        return header, rows
    raw = p.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "big5", "cp950", "gb18030", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    header = None
    rows = []
    for r in reader:
        if header is None:
            if any(c.strip() for c in r):
                header = [c.strip() or f"Column {i+1}" for i, c in enumerate(r)]
            continue
        if not any(c.strip() for c in r):
            continue
        rows.append({header[i]: r[i] for i in range(min(len(header), len(r)))})
        if max_rows and len(rows) >= max_rows:
            break
    return header or [], rows


def _h(s: str) -> str:
    return re.sub(r"[\s_\-()（）]+", " ", str(s).lower()).strip()


def suggest_mapping(kind: str, columns: list[str]) -> dict[str, str | None]:
    if kind == "COMPETITOR_CRAWL":
        bc_cols = breadcrumb_columns(columns)
        out = _suggest(FIELDS[kind], [c for c in columns if c not in bc_cols])
        if bc_cols and not out.get("breadcrumb"):
            out["_breadcrumb_columns"] = "|".join(bc_cols)  # joined in order at import
        return out
    return _suggest(FIELDS[kind], columns)


def _suggest(fields, columns: list[str]) -> dict[str, str | None]:
    used: set[str] = set()
    out: dict[str, str | None] = {}
    norm_cols = {c: _h(c) for c in columns}
    for f, (label, _req, aliases) in fields.items():
        choice = None
        for alias in [_h(label)] + [_h(a) for a in aliases]:
            choice = next((c for c, n in norm_cols.items() if n == alias and c not in used), None)
            if choice:
                break
        out[f] = choice
        if choice:
            used.add(choice)
    return out


def validate_mapping(kind: str, mapping: dict[str, str | None]) -> list[str]:
    errs = []
    if kind == "COMPETITOR_CRAWL":
        if not mapping.get("url"):
            errs.append("URL is required")
        if not (mapping.get("breadcrumb") or mapping.get("_breadcrumb_columns") or mapping.get("h1") or mapping.get("title")):
            errs.append("Map a Breadcrumb column (or at least H1 / Title)")
        return errs
    if kind == "TAXONOMY":
        if not mapping.get("url"):
            errs.append("Category URL is required")
        if not (mapping.get("name_cn") or mapping.get("h1_cn") or mapping.get("name_en")):
            errs.append("Category Name or H1 is required")
        if not (mapping.get("breadcrumb_cn") or mapping.get("parent_external_id")):
            errs.append("Breadcrumb is required (or Parent ID to rebuild it)")
    else:
        for f, (label, req, _) in FIELDS[kind].items():
            if req and not mapping.get(f):
                errs.append(f"{label} is required")
    return errs


def _s(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s or None


def _i(v) -> int | None:
    if v is None or str(v).strip() == "":
        return None
    try:
        return int(float(str(v).replace(",", "").strip()))
    except ValueError:
        return None


def import_rows(db: Session, kind: str, project_id: int, path: str, mapping: dict[str, str | None], replace: bool = True) -> dict:
    _, rows = read_table(path)
    get = lambda r, f: r.get(mapping[f]) if mapping.get(f) else None  # noqa: E731
    skipped = 0
    if kind == "TAXONOMY":
        if replace:
            db.execute(delete(Category).where(Category.project_id == project_id))
        for r in rows:
            name = _s(get(r, "name_cn")) or _s(get(r, "h1_cn")) or _s(get(r, "name_en"))
            if not name and not _s(get(r, "breadcrumb_cn")):
                skipped += 1
                continue
            db.add(Category(
                project_id=project_id,
                external_id=_s(get(r, "external_id")),
                url=_s(get(r, "url")),
                name_cn=_s(get(r, "name_cn")) or _s(get(r, "h1_cn")),
                name_en=_s(get(r, "name_en")) or _s(get(r, "h1_en")),
                h1_cn=_s(get(r, "h1_cn")),
                h1_en=_s(get(r, "h1_en")),
                breadcrumb_cn=_s(get(r, "breadcrumb_cn")),
                breadcrumb_en=_s(get(r, "breadcrumb_en")),
                parent_ref=_s(get(r, "parent_ref")),
                parent_external_id=_s(get(r, "parent_external_id")),
                product_count=_i(get(r, "product_count")),
                status=_s(get(r, "status")),
                raw={k: (str(v) if v is not None else None) for k, v in r.items()},
            ))
    elif kind == "SRP":
        if replace:
            db.execute(delete(SRP).where(SRP.project_id == project_id))
        seen: set[str] = set()
        for r in rows:
            kw = _s(get(r, "keyword"))
            if not kw:
                skipped += 1
                continue
            k = norm_key(kw)
            if k in seen:
                skipped += 1
                continue
            seen.add(k)
            db.add(SRP(
                project_id=project_id,
                keyword=kw,
                keyword_norm=k,
                url=_s(get(r, "url")),
                h1=_s(get(r, "h1")),
                keyword_en=_s(get(r, "keyword_en")),
                search_volume=_i(get(r, "search_volume")),
                product_count=_i(get(r, "product_count")),
                status=_s(get(r, "status")),
                raw={k2: (str(v) if v is not None else None) for k2, v in r.items()},
            ))
    elif kind == "SYNONYMS":
        existing = {s.alias_norm: s for s in db.query(Synonym).filter(Synonym.project_id == project_id)}
        for r in rows:
            alias, canon = _s(get(r, "alias")), _s(get(r, "canonical"))
            if not alias or not canon:
                skipped += 1
                continue
            k = norm_key(alias)
            s = existing.get(k) or Synonym(project_id=project_id, alias=alias, alias_norm=k)
            s.canonical = canon
            s.preferred_category_name = _s(get(r, "preferred_category_name"))
            s.preferred_node_id = _i(get(r, "preferred_node_id"))
            s.notes = _s(get(r, "notes"))
            s.approved_by = _s(get(r, "approved_by")) or "import"
            s.source = "HUMAN"
            if s.id is None:
                db.add(s)
                existing[k] = s
    db.flush()
    return {"rows": len(rows), "imported": len(rows) - skipped, "skipped": skipped}
