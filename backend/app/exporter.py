"""Excel/CSV exports for every stage. XLSX: bold header, autofilter, frozen first row.
CSV: UTF-8 with BOM (opens correctly in Excel with Chinese text), first sheet."""
from __future__ import annotations

import csv
import io
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from .config import settings
from .models import (
    SRP, BreadcrumbAudit, CategoryNode, Competitor, CompetitorCategory, CompetitorMapping, CrawlJob, CrawlPage,
    Export, GapChild, Project, TaxonomyGap,
)


@dataclass
class Sheet:
    name: str
    columns: list[str]
    rows: list[list[Any]]


def _j(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (list, tuple)):
        return "; ".join(_j(x) for x in v)
    if isinstance(v, dict):
        return "; ".join(f"{k}: {_j(x)}" for k, x in v.items())
    return str(v)


def _nodes(db, pid) -> dict[int, CategoryNode]:
    return {n.id: n for n in db.scalars(select(CategoryNode).where(CategoryNode.project_id == pid))}


def _comp_names(db, pid) -> dict[int, str]:
    return {c.id: c.name for c in db.scalars(select(Competitor).where(Competitor.project_id == pid))}


# ------------------------------------------------------------------ 1
def taxonomy_tree(db: Session, pid: int) -> list[Sheet]:
    nodes = sorted(_nodes(db, pid).values(), key=lambda n: n.breadcrumb_cn)
    cols = ["Node ID", "Category ID", "Category URL", "Name CN", "Name EN", "Parent", "Depth", "Children", "Descendants",
            "Breadcrumb CN", "Breadcrumb EN", "Product Count", "Status", "Inferred From Breadcrumb"]
    rows = [[n.id, n.external_id, n.url, n.name_cn, n.name_en, n.parent_name, n.depth, n.children_count, n.descendant_count,
             n.breadcrumb_cn, n.breadcrumb_en, n.product_count, n.status, "YES" if n.is_inferred else ""] for n in nodes]
    tree = [["    " * n.depth + ("└── " if n.depth else "") + n.name_cn, n.name_en, n.depth, n.url] for n in nodes]
    return [Sheet("Taxonomy Tree", cols, rows), Sheet("Tree View", ["Hierarchy", "Name EN", "Depth", "URL"], tree)]


# ------------------------------------------------------------------ 2
AUDIT_COLS = ["Category", "Category URL", "Current Breadcrumb CN", "Current Breadcrumb EN", "Current Parent", "Issues", "Action",
              "Recommended Parent", "Recommended Parent Status", "Recommended Breadcrumb CN", "Recommended Breadcrumb EN",
              "Current Score", "Recommended Score", "Current Ancestors", "Recommended Ancestors", "Links Removed", "Links Added",
              "Internal Linking Improvement", "Internal Linking Explanation", "Reason", "Confidence", "Evidence", "Review Status"]


def _audit_row(a: BreadcrumbAudit) -> list:
    return [a.category_name, a.category_url, a.current_breadcrumb_cn, a.current_breadcrumb_en, a.current_parent, _j(a.issues), a.action,
            a.recommended_parent, a.recommended_parent_status, a.recommended_breadcrumb_cn, a.recommended_breadcrumb_en,
            a.current_score, a.recommended_score, " > ".join(a.current_ancestors or []), " > ".join(a.recommended_ancestors or []),
            _j([x["label"] + (" (irrelevant)" if x.get("irrelevant") else "") for x in a.links_removed or []]),
            _j([x["label"] + (" (new page)" if x.get("new_page") else "") for x in a.links_added or []]),
            a.linking_improvement, a.linking_explanation, a.reason, a.confidence, _j(a.evidence), a.review_status]


def breadcrumb_audit(db: Session, pid: int) -> list[Sheet]:
    audits = db.scalars(select(BreadcrumbAudit).where(BreadcrumbAudit.project_id == pid).order_by(BreadcrumbAudit.current_breadcrumb_cn)).all()
    score_cols = ["Category", "Breadcrumb", "Score Type"] + [
        "parent_semantic_relevance", "hierarchy_specificity", "duplicate_level_penalty", "missing_parent_penalty",
        "depth_quality", "srp_support", "competitor_support", "internal_linking_relevance", "Total"]
    score_rows = []
    for a in audits:
        for t in ("current", "recommended"):
            b = (a.score_breakdown or {}).get(t) or {}
            score_rows.append([a.category_name, a.current_breadcrumb_cn if t == "current" else a.recommended_breadcrumb_cn, t.title()]
                              + [b.get(k) for k in score_cols[3:-1]] + [sum(v for v in b.values() if isinstance(v, (int, float)))])
    return [Sheet("Breadcrumb Audit", AUDIT_COLS, [_audit_row(a) for a in audits]),
            Sheet("Quality Score Breakdown", score_cols, score_rows)]


# ------------------------------------------------------------------ 3/4/13
SRP_COLS = ["SRP Keyword", "SRP URL", "H1", "English Keyword", "English Source", "Search Volume", "Product Count", "Core Product",
            "Modifiers", "Canonical Concept", "Matched Category", "Matched Category URL", "Matched Category Breadcrumb",
            "Suggested SRP Breadcrumb CN", "Suggested SRP Breadcrumb EN", "Match Method", "Score", "Confidence",
            "Confidence Band", "Duplicate Category Intent", "Recommended Action", "Potential Taxonomy Gap", "Reason", "Review Status"]


def _srp_row(s: SRP, nodes) -> list:
    n = nodes.get(s.matched_node_id)
    return [s.keyword, s.url, s.h1, s.keyword_en or s.keyword_en_generated, s.en_source, s.search_volume, s.product_count,
            s.core_product, _j(s.modifiers), s.canonical_concept, n.name_cn if n else "NO_MATCH", n.url if n else None,
            n.breadcrumb_cn if n else None, s.breadcrumb_cn, s.breadcrumb_en, s.match_method, s.score, s.confidence,
            s.confidence_band, s.duplicate_intent, s.recommended_action, "YES" if s.potential_gap else "NO", s.reason, s.review_status]


def srp_mapping(db: Session, pid: int) -> list[Sheet]:
    nodes = _nodes(db, pid)
    srps = db.scalars(select(SRP).where(SRP.project_id == pid).order_by(SRP.search_volume.desc().nullslast())).all()
    return [Sheet("SRP Breadcrumb Mapping", SRP_COLS, [_srp_row(s, nodes) for s in srps])]


def srp_review_queue(db: Session, pid: int) -> list[Sheet]:
    nodes = _nodes(db, pid)
    srps = db.scalars(select(SRP).where(SRP.project_id == pid, SRP.review_status.in_(("NEEDS_REVIEW", "LOW_CONFIDENCE", "NO_MATCH")))
                      .order_by(SRP.search_volume.desc().nullslast())).all()
    cols = SRP_COLS + ["Candidate 1", "Candidate 2", "Candidate 3", "Reviewer Decision", "Reviewer Category", "Reviewer Notes"]
    rows = []
    for s in srps:
        c = [f"{x['name']} ({x['score']}) — {x['breadcrumb']}" for x in (s.candidates or [])[:3]]
        rows.append(_srp_row(s, nodes) + c + [None] * (3 - len(c)) + [None, None, None])
    return [Sheet("SRP Review Queue", cols, rows)]


def final_srp_mapping(db: Session, pid: int) -> list[Sheet]:
    nodes = _nodes(db, pid)
    srps = db.scalars(select(SRP).where(SRP.project_id == pid, SRP.review_status.in_(("APPROVED", "AUTO_APPROVED")))
                      .order_by(SRP.search_volume.desc().nullslast())).all()
    cols = ["SRP Keyword", "SRP URL", "H1", "English Keyword", "Search Volume", "Final Category", "Final Category URL",
            "Final SRP Breadcrumb CN", "Final SRP Breadcrumb EN", "Duplicate Category Intent", "Indexing Recommendation",
            "Approval", "Match Method", "Confidence"]
    rows = []
    for s in srps:
        n = nodes.get(s.matched_node_id)
        if not n:
            continue
        idx = "Do not index separately (duplicate of category)" if s.duplicate_intent == "YES" else ("Review canonical" if s.duplicate_intent == "POTENTIAL" else "Indexable SRP")
        rows.append([s.keyword, s.url, s.h1, s.keyword_en or s.keyword_en_generated, s.search_volume, n.name_cn, n.url,
                     s.breadcrumb_cn, s.breadcrumb_en, s.duplicate_intent, idx,
                     "Human approved" if s.review_status == "APPROVED" else "Auto approved (≥90)", s.match_method, s.confidence])
    return [Sheet("Final Approved SRP Mapping", cols, rows)]


# ------------------------------------------------------------------ 5/6
def _latest_crawl_pages(db, pid) -> list[CrawlPage]:
    jobs = db.scalars(select(CrawlJob).where(CrawlJob.project_id == pid).order_by(CrawlJob.id.desc())).all()
    latest: dict[int, int] = {}
    for j in jobs:
        latest.setdefault(j.competitor_id, j.id)
    if not latest:
        return []
    return db.scalars(select(CrawlPage).where(CrawlPage.crawl_job_id.in_(latest.values())).order_by(CrawlPage.competitor_id, CrawlPage.id)).all()


CRAWL_COLS = ["Competitor", "URL", "Status Code", "Canonical", "Indexability", "Title", "H1", "Breadcrumb", "Depth", "Parent URL",
              "Anchor Text From Parent", "Internal Links", "Child Category Links", "Child Category Link Count", "Product Links",
              "Page Type", "Category Confidence", "Category Name", "Error", "Crawled At"]


def _crawl_row(p: CrawlPage, comps) -> list:
    return [comps.get(p.competitor_id), p.url, p.status_code, p.canonical, p.indexability, p.title, p.h1, p.breadcrumb_text, p.depth,
            p.parent_url, p.anchor_text, p.internal_links_count, _j([f"{c['text']} ({c['url']})" for c in (p.child_category_links or [])[:30]]),
            len(p.child_category_links or []), p.product_link_count, p.page_type, p.category_confidence, p.category_name, p.error,
            p.crawled_at.strftime("%Y-%m-%d %H:%M") if p.crawled_at else None]


def crawl_raw(db: Session, pid: int) -> list[Sheet]:
    comps = _comp_names(db, pid)
    return [Sheet("Competitor Crawl Raw", CRAWL_COLS, [_crawl_row(p, comps) for p in _latest_crawl_pages(db, pid)])]


def crawl_clean(db: Session, pid: int) -> list[Sheet]:
    comps = _comp_names(db, pid)
    pages = [p for p in _latest_crawl_pages(db, pid) if p.page_type in ("CATEGORY", "SUBCATEGORY") and (p.status_code or 0) < 400]
    return [Sheet("Clean Categories", CRAWL_COLS, [_crawl_row(p, comps) for p in pages])]


# ------------------------------------------------------------------ 7/8
def competitor_tree(db: Session, pid: int) -> list[Sheet]:
    comps = _comp_names(db, pid)
    cats = db.scalars(select(CompetitorCategory).where(CompetitorCategory.project_id == pid)
                      .order_by(CompetitorCategory.competitor_id, CompetitorCategory.breadcrumb)).all()
    cols = ["Competitor", "Competitor Category", "Parent", "Children", "Depth", "URL", "Breadcrumb", "Hierarchy Source", "Inferred", "Product Links"]
    rows = [[comps.get(c.competitor_id), c.name, c.parent_name, _j(c.children_names), c.depth, c.url, c.breadcrumb,
             c.hierarchy_source, "YES" if c.is_inferred else "", c.product_link_count] for c in cats]
    view = [[comps.get(c.competitor_id), "    " * c.depth + ("└── " if c.depth else "") + c.name, c.url] for c in cats]
    return [Sheet("Competitor Taxonomy Tree", cols, rows), Sheet("Tree View", ["Competitor", "Hierarchy", "URL"], view)]


def competitor_vs_current(db: Session, pid: int) -> list[Sheet]:
    comps, nodes = _comp_names(db, pid), _nodes(db, pid)
    cats = {c.id: c for c in db.scalars(select(CompetitorCategory).where(CompetitorCategory.project_id == pid))}
    maps = db.scalars(select(CompetitorMapping).where(CompetitorMapping.project_id == pid)).all()
    cols = ["Competitor", "Competitor Category", "Competitor Breadcrumb", "Competitor URL", "Classification", "My Category",
            "My Category URL", "My Breadcrumb", "Competitor Parent Maps To", "Score", "Confidence", "Method", "Reason", "Review Status"]
    rows = []
    for m in sorted(maps, key=lambda m: (m.competitor_id, cats[m.competitor_category_id].breadcrumb if m.competitor_category_id in cats else "")):
        cc = cats.get(m.competitor_category_id)
        n = nodes.get(m.matched_node_id)
        pn = nodes.get(m.parent_matched_node_id)
        rows.append([comps.get(m.competitor_id), cc.name if cc else None, cc.breadcrumb if cc else None, cc.url if cc else None,
                     m.match_type, n.name_cn if n else None, n.url if n else None, n.breadcrumb_cn if n else None,
                     pn.name_cn if pn else None, m.score, m.confidence, m.method, m.reason, m.review_status])
    summary = defaultdict(lambda: defaultdict(int))
    for m in maps:
        summary[comps.get(m.competitor_id)][m.match_type] += 1
    types = ["EXACT MATCH", "SEMANTIC MATCH", "COMPETITOR MORE SPECIFIC", "COMPETITOR MORE GENERAL", "MISSING FROM MY TAXONOMY", "EXISTING BUT MISPLACED", "AMBIGUOUS"]
    srows = [[c] + [summary[c][t] for t in types] for c in summary]
    return [Sheet("Competitor vs Current", cols, rows), Sheet("Summary", ["Competitor"] + types, srows)]


# ------------------------------------------------------------------ 9/10
def _gaps(db, pid):
    gaps = db.scalars(select(TaxonomyGap).where(TaxonomyGap.project_id == pid)).all()
    children = defaultdict(list)
    if gaps:
        for c in db.scalars(select(GapChild).where(GapChild.gap_id.in_([g.id for g in gaps]))):
            children[c.gap_id].append(c)
    order = {"HIGH": 0, "MEDIUM": 1, "LOW": 2}
    return sorted(gaps, key=lambda g: (order.get(g.priority, 3), -g.score)), children


def _evidence(g: TaxonomyGap, comp: str) -> str:
    ev = (g.competitor_evidence or {}).get(comp) or []
    return _j([e["breadcrumb"] + (f" ({e['url']})" if e.get("url") else "") for e in ev])


def gap_analysis(db: Session, pid: int) -> list[Sheet]:
    gaps, children = _gaps(db, pid)
    comps = sorted({c for g in gaps for c in (g.competitor_evidence or {})})
    cols = ["Gap Type", "Priority", "Score", "Confidence", "Recommended Action", "Suggested Parent CN", "Suggested Parent EN", "Parent Status",
            "Main Category CN", "Main Category EN", "Main Category Status", "Children", "Existing URL", "Current Breadcrumb",
            "Suggested Breadcrumb CN", "Suggested Breadcrumb EN", "EN Source", "Supporting SRPs", "SRP Count", "Combined Search Volume",
            "Product Count", "Competitor Count"] + [f"{c} Evidence" for c in comps] + [
            "Reason", "Links Added", "Links Removed", "Internal Linking Improvement", "Score Breakdown", "Review Status"]
    rows = []
    for g in gaps:
        rows.append([g.gap_type, g.priority, g.score, g.confidence, g.recommended_action, g.parent_name_cn, g.parent_name_en, g.parent_status,
                     g.main_name_cn, g.main_name_en, g.main_status, _j([f"{c.name_cn} [{c.status}]" for c in children[g.id]]),
                     g.existing_url, g.current_breadcrumb_cn, g.suggested_breadcrumb_cn, g.suggested_breadcrumb_en, g.en_source,
                     _j([f"{s['keyword']} ({s.get('search_volume') or 0})" for s in g.supporting_srps or []]), g.srp_count,
                     g.combined_search_volume, g.product_count, g.competitor_count]
                    + [_evidence(g, c) for c in comps]
                    + [g.reason, _j(g.links_added), _j(g.links_removed), g.linking_improvement, _j(g.score_breakdown), g.review_status])
    return [Sheet("Taxonomy Gap Analysis", cols, rows)]


def parent_child_gaps(db: Session, pid: int) -> list[Sheet]:
    gaps, children = _gaps(db, pid)
    comps = sorted({c for g in gaps for c in (g.competitor_evidence or {})})
    known = {"MOMO", "PCHOME"}
    named = [c for c in comps if c.upper() in known]
    other = [c for c in comps if c.upper() not in known]
    cols = ["Gap Type", "Suggested Parent CN", "Suggested Parent EN", "Parent Status", "Suggested Main Category CN",
            "Suggested Main Category EN", "Main Category Status", "Child Category CN", "Child Category EN", "Child Status",
            "Existing Child URL", "Existing Child Breadcrumb", "Suggested Breadcrumb CN", "Suggested Breadcrumb EN",
            "Supporting SRPs", "Combined Search Volume"] + [f"{c} Evidence" for c in named] + [
            "Other Competitor Evidence", "Reason", "Confidence", "Priority", "Recommended Action", "Review Status"]
    rows = []
    for g in gaps:
        if g.gap_type == "DUPLICATE" and not children[g.id]:
            continue
        kids = children[g.id] or [None]
        for c in kids:
            if c is None:  # single-node gap (missing child / misplaced / SRP-only)
                child = [g.main_name_cn, g.main_name_en, g.main_status, g.existing_url, g.current_breadcrumb_cn, g.suggested_breadcrumb_cn, g.suggested_breadcrumb_en]
                sup, vol = g.supporting_srps, g.combined_search_volume
            else:
                child = [c.name_cn, c.name_en, c.status, c.existing_url, c.existing_breadcrumb, c.suggested_breadcrumb_cn, c.suggested_breadcrumb_en]
                sup, vol = (c.supporting_srps or g.supporting_srps), (c.search_volume or g.combined_search_volume)
            rows.append([g.gap_type, g.parent_name_cn, g.parent_name_en, g.parent_status, g.main_name_cn, g.main_name_en, g.main_status]
                        + child + [_j([f"{s['keyword']} ({s.get('search_volume') or 0})" for s in sup or []]), vol]
                        + [_evidence(g, n) for n in named]
                        + ["; ".join(f"{o}: {_evidence(g, o)}" for o in other if _evidence(g, o)),
                           g.reason, g.confidence, g.priority, g.recommended_action, g.review_status])
    return [Sheet("Parent Child Gap Analysis", cols, rows)]


# ------------------------------------------------------------------ 11
def optimization_recommendations(db: Session, pid: int) -> list[Sheet]:
    audits = db.scalars(select(BreadcrumbAudit).where(BreadcrumbAudit.project_id == pid, BreadcrumbAudit.action != "KEEP")).all()
    gaps, children = _gaps(db, pid)
    cols = ["Source", "Recommendation Type", "Category / Page", "Current Breadcrumb CN", "Recommended Breadcrumb CN",
            "Recommended Breadcrumb EN", "Recommended Action", "Current Score", "Recommended Score", "Links Removed", "Links Added",
            "Internal Linking Improvement", "Internal Linking Explanation", "Reason", "Confidence", "Evidence", "Priority", "Review Status"]
    rows = []
    for a in audits:
        rows.append(["Breadcrumb Audit", a.action, a.category_name, a.current_breadcrumb_cn, a.recommended_breadcrumb_cn,
                     a.recommended_breadcrumb_en, a.action, a.current_score, a.recommended_score,
                     _j([x["label"] for x in a.links_removed or []]), _j([x["label"] for x in a.links_added or []]),
                     a.linking_improvement, a.linking_explanation, a.reason, a.confidence, _j(a.evidence),
                     "HIGH" if a.linking_improvement == "HIGH" else ("MEDIUM" if a.action != "REVIEW" else "LOW"), a.review_status])
    for g in gaps:
        if g.recommended_action and g.recommended_action.startswith(("MONITOR", "DO NOT")):
            continue
        rows.append(["Taxonomy Gap", g.gap_type, g.main_name_cn, g.current_breadcrumb_cn, g.suggested_breadcrumb_cn, g.suggested_breadcrumb_en,
                     g.recommended_action, None, None, _j(g.links_removed), _j(g.links_added), g.linking_improvement,
                     _j(g.links_added), g.reason, g.confidence,
                     f"{g.srp_count} SRPs / vol {g.combined_search_volume}; competitors: {', '.join((g.competitor_evidence or {}).keys())}",
                     g.priority, g.review_status])
    imp = {"HIGH": 0, "MEDIUM": 1, "LOW": 2, "NONE": 3, None: 4}
    rows.sort(key=lambda r: (imp.get(r[11], 4), r[1]))
    return [Sheet("Breadcrumb Optimization", cols, rows)]


# ------------------------------------------------------------------ 12
def final_taxonomy(db: Session, pid: int) -> list[Sheet]:
    nodes = _nodes(db, pid)
    root = next((n for n in nodes.values() if n.is_root), None)
    audits = {a.node_id: a for a in db.scalars(select(BreadcrumbAudit).where(BreadcrumbAudit.project_id == pid, BreadcrumbAudit.review_status == "APPROVED"))}
    gaps, children = _gaps(db, pid)
    gaps = [g for g in gaps if g.review_status == "APPROVED"]

    # virtual nodes: key -> {name_cn, name_en, parent(key), status, source, url, prev}
    V: dict[str, dict] = {}
    def nk(n: CategoryNode) -> str:
        return f"n{n.id}"
    for n in nodes.values():
        V[nk(n)] = {"name_cn": n.name_cn, "name_en": n.name_en, "parent": f"n{n.parent_id}" if n.parent_id else None,
                    "status": "UNCHANGED", "source": "", "url": n.url, "prev": n.breadcrumb_cn, "inferred": n.is_inferred, "root": n.is_root}
    def new_node(name, name_en, parent_key, source) -> str:
        key = f"new:{parent_key}:{name}"
        if key not in V:
            V[key] = {"name_cn": name, "name_en": name_en, "parent": parent_key, "status": "NEW", "source": source, "url": None,
                      "prev": None, "inferred": False, "root": False}
        return key

    for nid, a in audits.items():
        if nid not in nodes:
            continue
        k = nk(nodes[nid])
        src = f"Audit: {a.action}"
        if a.action == "KEEP":
            continue
        if a.action == "CONSOLIDATE":
            V[k]["status"], V[k]["source"] = f"CONSOLIDATE INTO {(a.evidence or {}).get('consolidate_into', '')}", src
            continue
        if a.action == "FLATTEN" and (a.evidence or {}).get("consolidate_into"):
            V[k]["status"], V[k]["source"] = f"CONSOLIDATE INTO {a.evidence['consolidate_into']}", src
            continue
        if a.recommended_parent_status == "NEW" and a.recommended_parent:
            anchor_labels = (a.recommended_ancestors or [])[:-1]
            anchor = next((n for n in nodes.values() if not n.is_root and n.breadcrumb_cn == " > ".join(([root.name_cn] if root else []) + anchor_labels)), root)
            pk = new_node(a.recommended_parent, None, nk(anchor) if anchor else None, src)
            V[k]["parent"], V[k]["status"], V[k]["source"] = pk, "MOVED", src
        elif a.recommended_parent_node_id and a.recommended_parent_node_id in nodes:
            V[k]["parent"], V[k]["status"], V[k]["source"] = f"n{a.recommended_parent_node_id}", ("FLATTENED" if a.action == "FLATTEN" else "MOVED"), src
        elif a.action == "FLATTEN":
            V[k]["status"], V[k]["source"] = "FLATTENED", src
            labels = a.recommended_ancestors or []
            par = next((n for n in nodes.values() if n.breadcrumb_cn == " > ".join(([root.name_cn] if root else []) + labels)), None)
            if par:
                V[k]["parent"] = nk(par)
        elif a.action == "CREATE CHILD":
            for d in (a.evidence or {}).get("child_demand", [])[:1]:
                new_node(d["name"], None, k, src)

    for g in gaps:
        src = f"Gap: {g.gap_type}"
        anchor = nodes.get(g.anchor_node_id) or nodes.get(g.parent_node_id)
        akey = nk(anchor) if anchor else (nk(root) if root else None)
        if g.gap_type in ("MISSING CHILD", "SRP-ONLY GAP"):
            new_node(g.main_name_cn, g.main_name_en, akey, src)
        elif g.gap_type in ("MISSING PARENT", "BOTH MISSING"):
            pk = new_node(g.parent_name_cn, g.parent_name_en, akey, src)
            for c in children[g.id]:
                if c.existing_node_id and c.existing_node_id in nodes:
                    ck = nk(nodes[c.existing_node_id])
                    V[ck]["parent"], V[ck]["status"], V[ck]["source"] = pk, "MOVED", src
                else:
                    new_node(c.name_cn, c.name_en, pk, src)
        elif g.gap_type == "MISPLACED EXISTING CATEGORY" and g.main_node_id in nodes and g.parent_node_id in nodes:
            ck = nk(nodes[g.main_node_id])
            V[ck]["parent"], V[ck]["status"], V[ck]["source"] = nk(nodes[g.parent_node_id]), "MOVED", src
        elif g.gap_type == "DUPLICATE":
            for c in children[g.id]:
                if c.existing_node_id in nodes:
                    ck = nk(nodes[c.existing_node_id])
                    V[ck]["status"], V[ck]["source"] = f"CONSOLIDATE INTO {g.existing_url or g.suggested_breadcrumb_cn}", src

    memo: dict[str, tuple[list[str], list[str | None]]] = {}
    def path(key: str, guard=()) -> tuple[list[str], list[str | None]]:
        if key in memo:
            return memo[key]
        v = V[key]
        if v["parent"] and v["parent"] in V and key not in guard:
            pc, pe = path(v["parent"], guard + (key,))
        else:
            pc, pe = [], []
        memo[key] = (pc + [v["name_cn"]], pe + [v["name_en"]])
        return memo[key]

    cols = ["Change Status", "Name CN", "Name EN", "Final Breadcrumb CN", "Final Breadcrumb EN", "Previous Breadcrumb CN", "URL", "Depth", "Source of Change"]
    rows = []
    for key, v in V.items():
        if v["root"]:
            continue
        pc, pe = path(key)
        final_cn = " > ".join(pc)
        status = v["status"]
        if status == "UNCHANGED" and v["prev"] and final_cn != v["prev"]:
            status = "PATH CHANGED (ancestor moved)"
        rows.append([status, v["name_cn"], v["name_en"], final_cn, " > ".join(pe) if all(pe) else None, v["prev"], v["url"], len(pc) - 1, v["source"]])
    rows.sort(key=lambda r: r[3])
    return [Sheet("Final Approved Taxonomy", cols, rows)]


def focus_sheets(res: dict) -> list[Sheet]:
    cols = ["Competitor", "Competitor Category", "Competitor Breadcrumb", "Competitor URL", "Status in Coupang", "Classification",
            "Coupang Category", "Coupang URL", "Coupang Breadcrumb CN", "Coupang Breadcrumb EN", "Suggested Parent",
            "Suggested Parent Status", "Suggested Breadcrumb CN", "Suggested Breadcrumb EN", "Gap Type", "Recommended Action",
            "Supporting SRPs", "SRP Count", "Search Volume", "Products Seen", "Hierarchy Source", "Confidence", "Reason"]
    rows = [[r["competitor"], r["competitor_category"], r["competitor_breadcrumb"], r["competitor_url"], r["status"], r["classification"],
             r["coupang_category"], r["coupang_url"], r["coupang_breadcrumb_cn"], r["coupang_breadcrumb_en"], r["suggested_parent"],
             r["suggested_parent_status"], r["suggested_breadcrumb_cn"], r["suggested_breadcrumb_en"], r["gap_type"], r["recommended_action"],
             _j([f"{s['keyword']} ({s.get('search_volume') or 0})" for s in r["supporting_srps"]]), r["srp_count"], r["search_volume"],
             r["products_seen"], r["hierarchy_source"], r["confidence"], r["reason"]] for r in res["rows"]]
    urls = [[u["competitor"], u["url"], u["page_type"], u["breadcrumb"], u["h1"], u["title"], u["status_code"]] for u in res["urls"]]
    coupang = [[c["category"], c["name_en"], c["breadcrumb_cn"], c["url"], _j(c["competitors_with_it"]) or "none"] for c in res["coupang_categories"]]
    return [
        Sheet("Category Focus", cols, rows),
        Sheet("Competitor URLs", ["Competitor", "URL", "Page Type", "Breadcrumb", "H1", "Title", "Status Code"], urls),
        Sheet("Coupang Categories", ["Coupang Category", "Name EN", "Breadcrumb CN", "URL", "Competitors With It"], coupang),
    ]


EXPORTS: dict[str, tuple[str, Callable[[Session, int], list[Sheet]]]] = {
    "current_taxonomy_tree": ("Current Taxonomy Tree", taxonomy_tree),
    "breadcrumb_audit": ("Breadcrumb Audit", breadcrumb_audit),
    "srp_breadcrumb_mapping": ("SRP Breadcrumb Mapping", srp_mapping),
    "srp_review_queue": ("SRP Review Queue", srp_review_queue),
    "competitor_crawl_raw": ("Competitor Crawl Raw", crawl_raw),
    "competitor_clean_categories": ("Competitor Clean Categories", crawl_clean),
    "competitor_taxonomy_tree": ("Competitor Taxonomy Tree", competitor_tree),
    "competitor_vs_current": ("Competitor vs Current Taxonomy", competitor_vs_current),
    "taxonomy_gap_analysis": ("Taxonomy Gap Analysis", gap_analysis),
    "parent_child_gap_analysis": ("Parent Child Gap Analysis", parent_child_gaps),
    "breadcrumb_optimization_recommendations": ("Breadcrumb Optimization Recommendations", optimization_recommendations),
    "final_approved_taxonomy": ("Final Approved Taxonomy", final_taxonomy),
    "final_approved_srp_mapping": ("Final Approved SRP Mapping", final_srp_mapping),
}

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _cell(v):
    if isinstance(v, str):
        v = _ILLEGAL.sub("", v)
        return v[:32000]
    return v


def to_xlsx(sheets: list[Sheet]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for sh in sheets:
        ws = wb.create_sheet(title=sh.name[:31])
        ws.append(sh.columns)
        for r in sh.rows:
            ws.append([_cell(v) for v in r])
        for i, col in enumerate(sh.columns, start=1):
            c = ws.cell(row=1, column=i)
            c.font, c.fill = HEADER_FONT, HEADER_FILL
            c.alignment = Alignment(vertical="center", wrap_text=True)
            sample = [len(str(col))] + [len(str(r[i - 1])) if i - 1 < len(r) and r[i - 1] is not None else 0 for r in sh.rows[:200]]
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(60, int(max(sample) * 1.15) + 2))
        ws.freeze_panes = "A2"
        if sh.columns:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(sh.columns))}{max(1, len(sh.rows) + 1)}"
        ws.row_dimensions[1].height = 30
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(sheets: list[Sheet]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    sh = sheets[0]
    w.writerow(sh.columns)
    for r in sh.rows:
        w.writerow(["" if v is None else v for v in r])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def build_export(db: Session, pid: int, key: str, fmt: str) -> tuple[str, bytes, int]:
    if key not in EXPORTS:
        raise KeyError(key)
    title, fn = EXPORTS[key]
    sheets = fn(db, pid)
    project = db.get(Project, pid)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M")
    safe = re.sub(r"[^\w\-]+", "_", project.name if project else str(pid))[:40]
    fname = f"{title} - {safe} - {stamp}.{fmt}"
    data = to_xlsx(sheets) if fmt == "xlsx" else to_csv(sheets)
    rows = len(sheets[0].rows) if sheets else 0
    db.add(Export(project_id=pid, export_key=key, format=fmt, filename=fname, row_count=rows))
    db.flush()
    return fname, data, rows


def write_export(db: Session, pid: int, key: str, fmt: str) -> tuple[Path, int]:
    fname, data, rows = build_export(db, pid, key, fmt)
    path = settings.data_dir / "exports" / fname
    path.write_bytes(data)
    return path, rows
